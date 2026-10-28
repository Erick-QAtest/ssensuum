import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app


@pytest.fixture()
def client():
    with TestClient(app) as test_client:
        yield test_client


def first_therapist(client) -> str:
    return client.get("/admin/rh/staff").json()[0]["id"]


def upload(client, therapist_id=None, name="contrato.pdf", content=b"%PDF-1.4 prueba", category="contrato", title=""):
    data = {"category": category, "title": title}
    if therapist_id:
        data["therapistId"] = therapist_id
    return client.post("/admin/rh/documents", data=data, files={"file": (name, content, "application/octet-stream")})


def test_staff_lists_every_therapist_with_profile(client):
    staff = client.get("/admin/rh/staff").json()
    assert {row["id"] for row in staff} >= {row["id"] for row in client.get("/therapists").json()}
    assert all("documents" in row and "active" in row for row in staff)


def test_update_profile_normalizes_and_validates(client):
    therapist_id = first_therapist(client)
    body = {"fullName": "  Jessy   Pérez ", "phone": "+52 33 1234 5678", "email": " Jessy@Mail.com ", "hiredOn": "2024-02-01", "emergencyContact": "Mamá 3311112222"}
    updated = client.put(f"/admin/rh/staff/{therapist_id}", json=body)
    assert updated.status_code == 200
    row = updated.json()
    assert row["fullName"] == "Jessy Pérez"
    assert row["phone"] == "3312345678"
    assert row["email"] == "jessy@mail.com"
    assert row["hiredOn"] == "2024-02-01"

    assert client.put(f"/admin/rh/staff/{therapist_id}", json={"phone": "123"}).status_code == 422
    assert client.put(f"/admin/rh/staff/{therapist_id}", json={"email": "nope"}).status_code == 422
    assert client.put("/admin/rh/staff/nadie", json={}).status_code == 404


def test_upload_download_and_delete_document(client):
    therapist_id = first_therapist(client)
    created = upload(client, therapist_id, name="../../Contrato Jessy.pdf")
    assert created.status_code == 201
    document = created.json()
    assert document["fileName"] == "Contrato Jessy.pdf"
    assert document["title"] == "Contrato Jessy"
    assert document["categoryLabel"] == "Contrato firmado"

    listed = client.get("/admin/rh/documents", params={"therapistId": therapist_id}).json()
    assert document["id"] in [row["id"] for row in listed]
    assert document["id"] not in [row["id"] for row in client.get("/admin/rh/documents", params={"therapistId": "general"}).json()]

    downloaded = client.get(f"/admin/rh/documents/{document['id']}/download")
    assert downloaded.status_code == 200
    assert downloaded.content == b"%PDF-1.4 prueba"
    assert "attachment" in downloaded.headers["content-disposition"]

    assert client.delete(f"/admin/rh/documents/{document['id']}").status_code == 204
    assert client.get(f"/admin/rh/documents/{document['id']}/download").status_code == 404


def test_rejects_invalid_uploads(client):
    assert upload(client, name="script.html").status_code == 422
    assert upload(client, content=b"").status_code == 422
    assert upload(client, category="secreto").status_code == 422
    assert upload(client, therapist_id="nadie").status_code == 404
    assert upload(client, content=b"x" * (8 * 1024 * 1024 + 1)).status_code == 422


def test_exports_staff_and_documents(client):
    therapist_id = first_therapist(client)
    upload(client, therapist_id, name="ine.png", content=b"png", category="identificacion")
    upload(client, name="reglamento.pdf", category="otro")
    upload(client, name="reglamento.pdf", category="otro")

    excel = client.get("/admin/rh/staff/export")
    assert excel.status_code == 200
    sheet = load_workbook(io.BytesIO(excel.content)).active
    assert sheet["A1"].value == "ID"
    assert therapist_id in [row[0].value for row in sheet.iter_rows(min_row=2)]

    archive = client.get("/admin/rh/documents/export")
    assert archive.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(archive.content)).namelist()
    assert any(name.endswith("identificacion/ine.png") for name in names)
    assert "General/otro/reglamento.pdf" in names
    assert "General/otro/reglamento (2).pdf" in names


def test_documentation_complete_when_required_documents_are_in(client):
    therapist_id = client.get("/admin/rh/staff").json()[-1]["id"]

    def row():
        return next(item for item in client.get("/admin/rh/staff").json() if item["id"] == therapist_id)

    assert row()["documentationComplete"] is False
    upload(client, therapist_id, name="contrato.pdf", category="contrato")
    upload(client, therapist_id, name="imagen.pdf", category="autorizacion_imagen")
    upload(client, therapist_id, name="descuento.pdf", category="autorizacion_descuento")
    assert [item["category"] for item in row()["missingDocuments"]] == ["identificacion"]
    upload(client, therapist_id, name="ine.jpg", category="identificacion")
    assert row()["documentationComplete"] is True
    assert row()["missingDocuments"] == []


def test_templates_live_apart_from_the_repository(client):
    template = upload(client, name="Contrato en blanco.pdf", category="plantilla").json()
    assert template["categoryLabel"] == "Formato para firmar"
    assert template["id"] in [row["id"] for row in client.get("/admin/rh/documents", params={"category": "plantilla"}).json()]
    assert template["id"] not in [row["id"] for row in client.get("/admin/rh/documents").json()]
    assert upload(client, first_therapist(client), category="plantilla").status_code == 422


JPEG = b"\xff\xd8\xff\xe0" + b"0" * 64


def hire(client, name="Nuevo Prueba", photo=JPEG, **fields):
    data = {"name": name, "strengths": "Calma, Técnica, Escucha", "bio": "Bio de prueba.", "fullName": "Nuevo Prueba López", "phone": "3312345678", **fields}
    return client.post("/admin/rh/staff", data=data, files={"photo": ("foto.jpg", photo, "image/jpeg")})


def test_hire_publishes_the_new_therapist_with_photo(client):
    created = hire(client, name="Ximena Alta")
    assert created.status_code == 201, created.text
    member = created.json()
    assert member["id"] == "ximena-alta" and member["active"] and member["isNew"]
    assert member["strengths"] == "Calma · Técnica · Escucha"
    assert member["fullName"] == "Nuevo Prueba López" and member["phone"] == "3312345678"
    assert member["photoUrl"].startswith("/therapists/ximena-alta/photo?v=")

    photo = client.get(member["photoUrl"])
    assert photo.status_code == 200 and photo.content == JPEG
    assert "ximena-alta" in {row["id"] for row in client.get("/therapists").json()}


def test_hire_validates_before_creating_anything(client):
    assert hire(client, name="Sin Foto", photo=b"no es imagen").status_code == 422
    assert hire(client, name="Mal Tel", phone="123").status_code == 422
    assert hire(client, name="  ").status_code == 422
    ids = {row["id"] for row in client.get("/admin/rh/staff").json()}
    assert not ids & {"sin-foto", "mal-tel"}
    # Un nombre que ya existe se reactiva, no se duplica.
    assert hire(client, name="jessy").status_code == 409


def test_photo_endpoint_rejects_odd_ids(client):
    assert client.get("/therapists/*/photo").status_code == 404
    assert client.get("/therapists/nadie/photo").status_code == 404


def test_replace_photo_changes_the_url(client):
    member = hire(client, name="Foto Cambio").json()
    png = b"\x89PNG\r\n\x1a\n" + b"1" * 32
    updated = client.put(f"/admin/rh/staff/{member['id']}/photo", files={"photo": ("f.png", png, "image/png")})
    assert updated.status_code == 200
    assert client.get(updated.json()["photoUrl"]).content == png


def test_leave_and_rehire(client):
    member = hire(client, name="Baja Prueba").json()
    check = client.get(f"/admin/rh/staff/{member['id']}/offboarding").json()
    assert check == {"upcomingReservations": [], "pendingPayouts": 0, "pendingAmount": 0}

    left = client.post(f"/admin/rh/staff/{member['id']}/leave", json={"leftOn": "2026-10-28", "reason": "Renuncia"})
    assert left.status_code == 200
    assert not left.json()["active"] and left.json()["leftOn"] == "2026-10-28" and left.json()["leaveReason"] == "Renuncia"
    assert member["id"] not in {row["id"] for row in client.get("/therapists").json()}
    assert client.post(f"/admin/rh/staff/{member['id']}/leave", json={"leftOn": "2026-10-28"}).status_code == 409

    back = client.post(f"/admin/rh/staff/{member['id']}/rehire").json()
    assert back["active"] and back["leftOn"] is None
    assert member["id"] in {row["id"] for row in client.get("/therapists").json()}


def test_public_profile_can_be_edited(client):
    member = hire(client, name="Perfil Edit").json()
    body = {"name": "Perfil Editado", "strengths": "Uno · Dos", "bio": "Nueva bio", "phone": member["phone"]}
    row = client.put(f"/admin/rh/staff/{member['id']}", json=body).json()
    assert row["name"] == "Perfil Editado" and row["strengths"] == "Uno · Dos" and row["bio"] == "Nueva bio"


def test_leave_is_blocked_while_there_are_upcoming_reservations(client):
    import json
    import uuid
    from datetime import datetime, timedelta

    from app.db import SessionLocal
    from app.domain.rules import TZ
    from app.models import Reservation

    member = hire(client, name="Con Agenda").json()
    starts = (datetime.now(TZ) + timedelta(days=5)).replace(hour=11, minute=0, second=0, microsecond=0)
    with SessionLocal() as db:
        db.add(Reservation(
            id=str(uuid.uuid4()), service_id="massage-special", therapist_ids_json=json.dumps([member["id"]]),
            starts_at=starts, ends_at=starts + timedelta(hours=1), status="CONFIRMED", deposit_amount=300, source="admin",
        ))
        db.commit()

    check = client.get(f"/admin/rh/staff/{member['id']}/offboarding").json()
    assert len(check["upcomingReservations"]) == 1
    blocked = client.post(f"/admin/rh/staff/{member['id']}/leave", json={"leftOn": "2026-10-28"})
    assert blocked.status_code == 409 and "1 reserva próxima" in blocked.json()["detail"]
    handled = client.post(f"/admin/rh/staff/{member['id']}/leave", json={"leftOn": "2026-10-28", "reservationsHandled": True})
    assert handled.status_code == 200 and not handled.json()["active"]
