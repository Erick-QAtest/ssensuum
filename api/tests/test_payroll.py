import io
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.domain.payroll import settle, week_start
from app.main import app


@pytest.fixture()
def client():
    with TestClient(app) as test_client:
        yield test_client


def test_week_runs_monday_to_sunday():
    from datetime import date
    assert week_start(date(2026, 10, 19)) == date(2026, 10, 19)  # lunes
    assert week_start(date(2026, 10, 25)) == date(2026, 10, 19)  # domingo


def test_cleaning_comes_off_the_first_service_only():
    result = settle([{"serviceName": "a"}, {"serviceName": "b"}, {"serviceName": "c"}], tips=250)
    assert [line["pay"] for line in result["lines"]] == [400, 500, 500]
    assert (result["gross"], result["cleaning"], result["servicePay"], result["total"]) == (1500, 100, 1400, 1650)


def test_no_services_no_cleaning():
    assert settle([], tips=0) == {"services": 0, "gross": 0, "cleaning": 0, "servicePay": 0, "tips": 0, "total": 0, "lines": []}


def record(client, therapist_ids, paid_at, tips=(), service_id="massage-special"):
    body = {
        "customerName": "Cliente Prueba",
        "customerPhone": "34" + str(uuid.uuid4().int)[:8],
        "serviceId": service_id,
        "therapistIds": therapist_ids,
        "paidAt": paid_at,
        "amount": 1500,
        "paymentMethod": "efectivo",
        "tips": [{"therapistId": tid, "amount": amount} for tid, amount in tips],
        "tipMethod": "efectivo",
    }
    response = client.post("/admin/records", json=body)
    assert response.status_code == 201
    return response.json()


def test_weekly_payroll_per_therapist(client):
    a, b = [row["id"] for row in client.get("/therapists").json()][:2]
    # Semana del lunes 2 al domingo 8 de febrero de 2026; el lunes 9 ya es otra semana.
    second = record(client, [a], "2026-02-04T12:00:00-06:00", tips=[(a, 200)])
    first = record(client, [a, b], "2026-02-02T09:30:00-06:00", service_id="four-hands")
    record(client, [b], "2026-02-08T23:30:00-06:00")
    record(client, [a], "2026-02-09T10:00:00-06:00")

    payroll = client.get("/admin/rh/payroll", params={"week": "2026-02-05"}).json()
    assert (payroll["weekStart"], payroll["weekEnd"]) == ("2026-02-02", "2026-02-08")
    rows = {row["therapistId"]: row for row in payroll["staff"]}

    assert rows[a]["services"] == 2
    assert [line["recordId"] for line in rows[a]["lines"]] == [first["id"], second["id"]]
    assert [line["pay"] for line in rows[a]["lines"]] == [400, 500]
    assert (rows[a]["servicePay"], rows[a]["tips"], rows[a]["total"]) == (900, 200, 1100)

    assert rows[b]["services"] == 2
    assert rows[b]["total"] == 900
    assert payroll["totals"]["total"] == 2000

    excel = client.get("/admin/rh/payroll/export", params={"week": "2026-02-05"})
    assert excel.status_code == 200
    book = load_workbook(io.BytesIO(excel.content))
    assert book["Liquidación"].cell(row=book["Liquidación"].max_row, column=7).value == 2000
    assert book["Detalle"].max_row == 5
