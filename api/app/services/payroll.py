import io
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from ..domain.payroll import settle, week_start
from ..domain.rules import TZ, ensure_tz
from ..repositories import hr as hr_repo
from ..repositories import records as records_repo


def weekly_payroll(db: Session, day: date) -> dict:
    monday = week_start(day)
    start = datetime.combine(monday, time.min, tzinfo=TZ)
    # El repositorio los da del más reciente al más antiguo; la limpieza va sobre el primero de la semana.
    rows = sorted(records_repo.list_paid_between(db, start, start + timedelta(days=7)), key=lambda row: row.paid_at)
    names = {row.id: row.name for row in hr_repo.list_therapists(db)}

    services: dict[str, list[dict]] = {}
    tips: dict[str, int] = {}
    for record in rows:
        for therapist_id in records_repo.therapist_ids(record):
            services.setdefault(therapist_id, []).append({
                "recordId": record.id,
                "paidAt": ensure_tz(record.paid_at).isoformat(),
                "serviceName": record.service_name,
                "customerName": record.customer_name,
            })
        for tip in records_repo.tips(record):
            tips[tip["therapistId"]] = tips.get(tip["therapistId"], 0) + tip["amount"]

    staff = [
        {"therapistId": therapist_id, "therapistName": names.get(therapist_id, therapist_id), **settle(services.get(therapist_id, []), tips.get(therapist_id, 0))}
        for therapist_id in services.keys() | tips.keys()
    ]
    staff.sort(key=lambda row: row["therapistName"])
    return {
        "weekStart": monday.isoformat(),
        "weekEnd": (monday + timedelta(days=6)).isoformat(),
        "staff": staff,
        "totals": {
            key: sum(row[key] for row in staff)
            for key in ("services", "gross", "cleaning", "servicePay", "tips", "total")
        },
    }


def export_payroll_xlsx(db: Session, day: date) -> bytes:
    payroll = weekly_payroll(db, day)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Liquidación"
    summary.append([f"Semana del {payroll['weekStart']} al {payroll['weekEnd']}"])
    summary.append(["Masajista", "Servicios", "Por servicios", "Limpieza", "Pago neto", "Propinas", "Total a pagar"])
    for row in payroll["staff"]:
        summary.append([row["therapistName"], row["services"], row["gross"], -row["cleaning"], row["servicePay"], row["tips"], row["total"]])
    totals = payroll["totals"]
    summary.append(["Total", totals["services"], totals["gross"], -totals["cleaning"], totals["servicePay"], totals["tips"], totals["total"]])

    detail = workbook.create_sheet("Detalle")
    detail.append(["Masajista", "Fecha", "Servicio", "Cliente", "Por servicio", "Limpieza", "Pago"])
    for row in payroll["staff"]:
        for line in row["lines"]:
            paid_at = datetime.fromisoformat(line["paidAt"]).strftime("%Y-%m-%d %H:%M")
            detail.append([row["therapistName"], paid_at, line["serviceName"], line["customerName"], line["gross"], -line["cleaning"], line["pay"]])

    for sheet, header_row in ((summary, 2), (detail, 1)):
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column[header_row - 1:])
            sheet.column_dimensions[column[0].column_letter].width = min(width + 2, 40)
    summary[1][0].font = Font(bold=True)
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
