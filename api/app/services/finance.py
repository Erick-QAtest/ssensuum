import io
import uuid
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from ..domain.errors import NotFoundError, ValidationError
from ..domain.finance import (
    CASH_ACCOUNTS,
    EXPENSE_CATEGORIES,
    FINANCE_START,
    MOVEMENT_KINDS,
    OTHER_INCOME_CATEGORIES,
    PAYMENT_METHOD_ACCOUNT,
    PAYOUT_ACCOUNT,
    Entry,
    balance_sheet,
    balances,
    cash_flow,
    income_statement,
    movement_entry,
)
from ..domain.rules import TZ, ensure_tz
from ..models import FinanceMovement
from ..repositories import finance as finance_repo
from ..repositories import records as records_repo
from ..schemas import FinanceMovementCreate
from . import payroll as payroll_service


def meta() -> dict:
    return {
        "startDate": FINANCE_START.isoformat(),
        "accounts": [{"value": code, "label": label} for code, label in CASH_ACCOUNTS.items()],
        "kinds": [{"value": code, "label": label} for code, (label, _) in MOVEMENT_KINDS.items()],
        "expenseCategories": [{"value": code, "label": label} for code, label in EXPENSE_CATEGORIES.items()],
        "incomeCategories": [{"value": code, "label": label} for code, label in OTHER_INCOME_CATEGORIES.items()],
    }


def _validate(payload: FinanceMovementCreate) -> None:
    if payload.kind not in MOVEMENT_KINDS:
        raise ValidationError("Tipo de movimiento no válido")
    if payload.account not in CASH_ACCOUNTS:
        raise ValidationError("Cuenta no válida")
    if not payload.description.strip():
        raise ValidationError("Escribe una descripción")
    if payload.kind == "gasto" and payload.category not in EXPENSE_CATEGORIES:
        raise ValidationError("Elige la categoría del gasto")
    if payload.kind == "ingreso" and payload.category not in OTHER_INCOME_CATEGORIES:
        raise ValidationError("Elige la categoría del ingreso")
    if payload.kind == "traspaso" and (payload.to_account not in CASH_ACCOUNTS or payload.to_account == payload.account):
        raise ValidationError("Elige una cuenta destino distinta a la de origen")


def create_movement(db: Session, payload: FinanceMovementCreate) -> FinanceMovement:
    _validate(payload)
    keeps_category = payload.kind in ("gasto", "ingreso")
    return finance_repo.save(db, FinanceMovement(
        id=str(uuid.uuid4()),
        day=payload.day.isoformat(),
        kind=payload.kind,
        category=payload.category if keeps_category else None,
        description=" ".join(payload.description.split()),
        amount=payload.amount,
        account=payload.account,
        to_account=payload.to_account if payload.kind == "traspaso" else None,
    ))


def delete_movement(db: Session, movement_id: str) -> None:
    movement = finance_repo.get(db, movement_id)
    if not movement:
        raise NotFoundError("Movimiento no encontrado")
    finance_repo.delete(db, movement)


def list_movements(db: Session, start: date, end: date) -> list[FinanceMovement]:
    return finance_repo.list_movements(db, start, end)


def journal(db: Session, until: date) -> list[Entry]:
    """Todos los asientos hasta esa fecha, desde el inicio de la contabilidad."""
    entries: list[Entry] = []
    if until >= FINANCE_START:
        start = datetime.combine(FINANCE_START, time.min, tzinfo=TZ)
        records = records_repo.list_paid_between(db, start, datetime.combine(until + timedelta(days=1), time.min, tzinfo=TZ))
        for record in records:
            day = ensure_tz(record.paid_at).date()
            memo = f"{record.service_name} · {record.customer_name}"
            entries.append(
                Entry(day, memo, ("operacion", "Cobros por servicios"))
                .debit(PAYMENT_METHOD_ACCOUNT[record.payment_method], record.amount)
                .credit("ingresos_servicios", record.amount)
            )
            tips = sum(tip["amount"] for tip in records_repo.tips(record))
            if tips:
                account = PAYMENT_METHOD_ACCOUNT[record.tip_method or record.payment_method]
                entries.append(Entry(day, memo, ("operacion", "Propinas recibidas")).debit(account, tips).credit("propinas_por_pagar", tips))

        for line in payroll_service.lines_between(db, FINANCE_START, until):
            day = datetime.fromisoformat(line["paidAt"]).date()
            memo = f"{line['therapistName']} · {line['serviceName']}"
            # Los honorarios se devengan al concluir el servicio; la limpieza autorizada se compensa ahí mismo.
            entries.append(
                Entry(day, memo, ("operacion", "Honorarios devengados"))
                .debit("honorarios", line["gross"]).credit("honorarios_por_pagar", line["gross"])
            )
            if line["cleaning"]:
                entries.append(
                    Entry(day, memo, ("operacion", "Limpieza compensada"))
                    .debit("honorarios_por_pagar", line["cleaning"]).credit("ingreso_limpieza", line["cleaning"])
                )
            if line["payoutAt"]:
                paid_on = datetime.fromisoformat(line["payoutAt"]).astimezone(TZ).date()
                if paid_on <= until:
                    fee = line["gross"] - line["cleaning"]
                    entries.append(Entry(paid_on, memo, ("operacion", "Pagos de honorarios")).debit("honorarios_por_pagar", fee).credit(PAYOUT_ACCOUNT, fee))
                    if line["tip"]:
                        entries.append(Entry(paid_on, memo, ("operacion", "Propinas entregadas")).debit("propinas_por_pagar", line["tip"]).credit(PAYOUT_ACCOUNT, line["tip"]))

    for movement in finance_repo.list_movements(db, end=until):
        entries.append(movement_entry(
            date.fromisoformat(movement.day), movement.kind, movement.amount, movement.account,
            movement.category, movement.to_account, movement.description,
        ))
    return entries


def statements(db: Session, start: date, end: date) -> dict:
    if end < start:
        raise ValidationError("La fecha final debe ser posterior a la inicial")
    entries = journal(db, end)
    before = [entry for entry in entries if entry.day < start]
    during = [entry for entry in entries if entry.day >= start]
    return {
        "from": start.isoformat(),
        "to": end.isoformat(),
        "startDate": FINANCE_START.isoformat(),
        "incomeStatement": income_statement(during),
        "cashFlow": cash_flow(before, during),
        "balanceSheet": balance_sheet(entries),
    }


def _buckets(start: date, end: date) -> tuple[str, list[tuple[date, date]]]:
    """Divide el periodo en días (hasta un mes), semanas (hasta cuatro meses) o meses."""
    days = (end - start).days + 1
    if days <= 31:
        return "day", [(start + timedelta(days=i), start + timedelta(days=i)) for i in range(days)]
    if days <= 124:
        result, cursor = [], start
        while cursor <= end:
            last = min(cursor + timedelta(days=6 - cursor.weekday()), end)
            result.append((cursor, last))
            cursor = last + timedelta(days=1)
        return "week", result
    result, cursor = [], start
    while cursor <= end:
        next_month = (cursor.replace(day=1) + timedelta(days=32)).replace(day=1)
        last = min(next_month - timedelta(days=1), end)
        result.append((cursor, last))
        cursor = last + timedelta(days=1)
    return "month", result


def _cash(entries: list[Entry]) -> int:
    totals = balances(entries)
    return sum(totals.get(code, 0) for code in CASH_ACCOUNTS)


def _payables(entries: list[Entry]) -> int:
    totals = balances(entries)
    return totals.get("honorarios_por_pagar", 0) + totals.get("propinas_por_pagar", 0)


def _top(rows: dict[str, dict], limit: int = 6) -> list[dict]:
    """Los conceptos más grandes y el resto sumado en "Otros"."""
    ordered = sorted(rows.values(), key=lambda row: -row["amount"])
    if len(ordered) <= limit:
        return ordered
    rest = ordered[limit - 1:]
    return ordered[: limit - 1] + [{"label": "Otros", "amount": sum(r["amount"] for r in rest), "count": sum(r.get("count", 0) for r in rest)}]


def dashboard(db: Session, start: date, end: date) -> dict:
    """Indicadores, series por periodo y desgloses para el tablero de Finanzas."""
    if end < start:
        raise ValidationError("La fecha final debe ser posterior a la inicial")
    length = (end - start).days + 1
    previous_start, previous_end = start - timedelta(days=length), start - timedelta(days=1)
    entries = journal(db, end)
    within = lambda first, last: [entry for entry in entries if first <= entry.day <= last]
    upto = lambda last: [entry for entry in entries if entry.day <= last]

    current = income_statement(within(start, end))
    previous = income_statement(within(previous_start, previous_end))

    # Antes del inicio de la contabilidad no hay servicios: la serie arranca ahí para no mostrar días vacíos.
    granularity, ranges = _buckets(max(start, FINANCE_START) if end >= FINANCE_START else start, end)
    series = []
    for first, last in ranges:
        pnl = income_statement(within(first, last))
        series.append({
            "start": first.isoformat(),
            "end": last.isoformat(),
            "revenue": pnl["revenue"],
            "costOfServices": pnl["costOfServices"],
            "operatingExpenses": pnl["operatingExpenses"],
            "outflows": pnl["costOfServices"] + pnl["operatingExpenses"],
            "otherIncome": pnl["otherIncomeTotal"],
            "netIncome": pnl["netIncome"],
            "cash": _cash(upto(last)),
        })

    costs = {"honorarios": {"label": "Honorarios de masajistas", "amount": current["costOfServices"]}} if current["costOfServices"] else {}
    for row in current["expenses"]:
        costs[row["account"]] = {"label": row["label"], "amount": row["amount"]}

    services: dict[str, dict] = {}
    methods: dict[str, dict] = {}
    if end >= FINANCE_START:
        first = max(start, FINANCE_START)
        records = records_repo.list_paid_between(
            db, datetime.combine(first, time.min, tzinfo=TZ), datetime.combine(end + timedelta(days=1), time.min, tzinfo=TZ),
        )
        for record in records:
            service = services.setdefault(record.service_name, {"label": record.service_name, "amount": 0, "count": 0})
            service["amount"] += record.amount
            service["count"] += 1
            account = PAYMENT_METHOD_ACCOUNT[record.payment_method]
            method = methods.setdefault(account, {"label": CASH_ACCOUNTS[account], "amount": 0, "count": 0})
            method["amount"] += record.amount
            method["count"] += 1

    return {
        "from": start.isoformat(),
        "to": end.isoformat(),
        "previousFrom": previous_start.isoformat(),
        "previousTo": previous_end.isoformat(),
        "granularity": granularity,
        "kpis": {
            "revenue": current["revenue"],
            "netIncome": current["netIncome"],
            "netMargin": current["netMargin"],
            "cash": _cash(upto(end)),
            "payables": _payables(upto(end)),
            "services": sum(row["count"] for row in services.values()),
            "previous": {
                "revenue": previous["revenue"],
                "netIncome": previous["netIncome"],
                "cash": _cash(upto(previous_end)),
                "payables": _payables(upto(previous_end)),
            },
        },
        "series": series,
        "costBreakdown": _top(costs, limit=7),
        "revenueByService": _top(services),
        "revenueByAccount": sorted(methods.values(), key=lambda row: -row["amount"]),
    }


def export_xlsx(db: Session, start: date, end: date) -> bytes:
    data = statements(db, start, end)
    period = f"Del {data['from']} al {data['to']}"
    workbook = Workbook()

    results = workbook.active
    results.title = "Estado de resultados"
    pnl = data["incomeStatement"]
    rows = [
        ("Estado de resultados", None), (period, None), (None, None),
        ("Ingresos por servicios", pnl["revenue"]),
        ("Honorarios de masajistas", -pnl["costOfServices"]),
        ("Utilidad bruta", pnl["grossProfit"]), (None, None), ("Gastos de operación", None),
        *[(f"  {row['label']}", -row["amount"]) for row in pnl["expenses"]],
        ("Total gastos de operación", -pnl["operatingExpenses"]),
        ("Utilidad de operación", pnl["operatingIncome"]),
        *([(None, None), ("Otros ingresos", None)] if pnl["otherIncome"] else []),
        *[(f"  {row['label']}", row["amount"]) for row in pnl["otherIncome"]],
        ("Utilidad neta", pnl["netIncome"]),
    ]
    for row in rows:
        results.append(row)

    flow_sheet = workbook.create_sheet("Flujo de efectivo")
    flow = data["cashFlow"]
    flow_sheet.append(("Flujo de efectivo",))
    flow_sheet.append((period,))
    flow_sheet.append(())
    flow_sheet.append(("Efectivo al inicio", flow["opening"]))
    for section in flow["sections"]:
        flow_sheet.append((section["label"],))
        for item in section["items"]:
            flow_sheet.append((f"  {item['label']}", item["amount"]))
        flow_sheet.append((f"Total {section['label'].lower()}", section["total"]))
    flow_sheet.append(("Cambio neto del periodo", flow["netChange"]))
    flow_sheet.append(("Efectivo al final", flow["closing"]))
    flow_sheet.append(())
    flow_sheet.append(("Cuenta", "Saldo inicial", "Saldo final"))
    for account in flow["accounts"]:
        flow_sheet.append((account["label"], account["opening"], account["closing"]))

    balance = workbook.create_sheet("Balance general")
    sheet = data["balanceSheet"]
    balance.append(("Balance general",))
    balance.append((f"Al {data['to']}",))
    balance.append(())
    for title, group, total in (
        ("Activo", sheet["assets"], sheet["totalAssets"]),
        ("Pasivo", sheet["liabilities"], sheet["totalLiabilities"]),
        ("Capital", sheet["equity"], sheet["totalEquity"]),
    ):
        balance.append((title,))
        for row in group:
            balance.append((f"  {row['label']}", row["amount"]))
        balance.append((f"Total {title.lower()}", total))
    balance.append(("Pasivo más capital", sheet["totalLiabilities"] + sheet["totalEquity"]))

    movements = workbook.create_sheet("Movimientos")
    movements.append(("Fecha", "Tipo", "Categoría", "Descripción", "Cuenta", "Cuenta destino", "Importe"))
    categories = {**EXPENSE_CATEGORIES, **OTHER_INCOME_CATEGORIES}
    for movement in sorted(list_movements(db, start, end), key=lambda row: row.day):
        movements.append((
            movement.day, MOVEMENT_KINDS[movement.kind][0], categories.get(movement.category or "", None),
            movement.description, CASH_ACCOUNTS[movement.account], CASH_ACCOUNTS.get(movement.to_account or ""), movement.amount,
        ))

    for worksheet in workbook.worksheets:
        worksheet.column_dimensions["A"].width = 44
        for column in "BCDEFG":
            worksheet.column_dimensions[column].width = 18
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, int):
                    cell.number_format = '"$"#,##0;-"$"#,##0'
                if isinstance(cell.value, str) and not cell.value.startswith("  ") and cell.column == 1:
                    cell.font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
