import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from typing import Any, Optional

from .errors import ValidationError

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000

# Encabezados aceptados por campo, ya normalizados (minúsculas, sin acentos ni signos).
COLUMN_ALIASES: dict[str, set[str]] = {
    "phone": {"telefono", "tel", "telefono celular", "numero de telefono", "numero", "celular", "cel", "movil", "whatsapp", "phone"},
    "first_name": {"nombre", "nombres", "nombre completo", "cliente", "name", "first name", "firstname"},
    "last_name": {"apellido", "apellidos", "last name", "lastname"},
    "email": {"correo", "correo electronico", "email", "e mail", "mail"},
    "birthday": {"cumpleanos", "cumple", "fecha de nacimiento", "fecha nacimiento", "nacimiento", "birthday"},
}
_BIRTHDAY_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S")


def normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def map_columns(headers: list[Any]) -> dict[str, int]:
    """Relaciona cada campo con el índice de su columna; gana la primera coincidencia."""
    columns: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalized = normalize_header(header)
        for field, aliases in COLUMN_ALIASES.items():
            if field not in columns and normalized in aliases:
                columns[field] = index
    return columns


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    # Excel guarda los teléfonos como número: 3310000001.0
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_birthday(value: Any) -> Optional[date]:
    """Devuelve la fecha o None si no se reconoce o no es creíble."""
    if isinstance(value, datetime):
        parsed: Optional[date] = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        parsed = None
        text = cell_text(value)
        for fmt in _BIRTHDAY_FORMATS:
            try:
                parsed = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
    if parsed is None or parsed.year < 1900 or parsed > date.today():
        return None
    return parsed


def _read_csv(content: bytes) -> list[list[Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    first_line = text.split("\n", 1)[0]
    delimiter = max(",;\t", key=first_line.count)
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError):
        raise ValidationError("No se pudo leer el archivo de Excel")
    try:
        return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_table(filename: str, content: bytes) -> tuple[list[Any], list[tuple[int, list[Any]]]]:
    """Devuelve los encabezados y las filas con datos como (número de fila en el archivo, celdas)."""
    if len(content) > MAX_FILE_BYTES:
        raise ValidationError("El archivo supera el límite de 5 MB")

    name = filename.lower()
    if name.endswith(".xlsx"):
        rows = _read_xlsx(content)
    elif name.endswith(".csv"):
        rows = _read_csv(content)
    else:
        raise ValidationError("Formato no soportado. Sube un archivo .xlsx o .csv")

    numbered = [(index, row) for index, row in enumerate(rows, start=1) if any(cell_text(cell) for cell in row)]
    if not numbered:
        raise ValidationError("El archivo está vacío")

    headers = numbered[0][1]
    data = numbered[1:]
    if len(data) > MAX_ROWS:
        raise ValidationError(f"El archivo supera el límite de {MAX_ROWS} filas")
    return headers, data
